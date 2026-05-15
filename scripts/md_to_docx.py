#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert docs/DOCUMENTATION_TECHNIQUE_SYSCO_FR.md to .docx
Requires: pip install python-docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_UNDERLINE
except ImportError:
    print("Missing dependency. Run:  pip install python-docx", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "DOCUMENTATION_TECHNIQUE_SYSCO_FR.md"
OUT_PATH = ROOT / "docs" / "DOCUMENTATION_TECHNIQUE_SYSCO_FR.docx"


def add_inline_runs(paragraph, text: str) -> None:
    """Split **bold** and `code` into runs."""
    pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+?`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        chunk = m.group(0)
        if chunk.startswith("**"):
            r = paragraph.add_run(chunk[2:-2])
            r.bold = True
        else:
            r = paragraph.add_run(chunk[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    j = start
    while j < len(lines) and is_table_row(lines[j]):
        parts = [c.strip() for c in lines[j].strip().split("|")]
        parts = [p for p in parts if p != ""]
        if parts and all(re.match(r"^:?-+$", re.sub(r"\s+", "", p)) for p in parts):
            j += 1
            continue
        rows.append(parts)
        j += 1
    return rows, j


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    try:
        table.style = "Table Grid"
    except (ValueError, KeyError):
        pass
    for ri, row in enumerate(rows):
        for ci in range(cols):
            cell = row[ci] if ci < len(row) else ""
            table.rows[ri].cells[ci].text = cell


def main() -> None:
    if not MD_PATH.is_file():
        print("Not found:", MD_PATH, file=sys.stderr)
        sys.exit(1)

    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("—" * 50)
            i += 1
            continue

        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue

        if is_table_row(line):
            rows, j = parse_table(lines, i)
            add_table(doc, rows)
            i = j
            continue

        if stripped.startswith(("- ", "* ")) and not stripped.startswith("**"):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, stripped[2:].strip())
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            r = p.add_run(stripped.lstrip(">").strip())
            r.italic = True
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line.rstrip())
        i += 1

    doc.save(OUT_PATH)
    print("OK:", OUT_PATH)


if __name__ == "__main__":
    main()
